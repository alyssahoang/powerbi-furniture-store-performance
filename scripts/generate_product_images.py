from __future__ import annotations

import argparse
import io
import re
import time
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable
from urllib.parse import quote

import openpyxl
import requests
from PIL import Image


API_BASE = "https://image.pollinations.ai/prompt/"
REFERENCE_STYLE = (
    "Match a premium furniture catalog contact-sheet style: single product only, "
    "isolated and centered on a seamless bright white background (#f8f7f7), "
    "high-key studio lighting, soft short grounding shadow directly beneath the item, "
    "warm natural oak or walnut wood tones, ivory, beige, taupe or soft grey fabrics when applicable, "
    "clean Scandinavian-Japandi retail aesthetic, crisp edges, realistic proportions, "
    "front or three-quarter front product angle when appropriate, square composition, "
    "no room scene, no wall texture, no floor horizon line, no text, no logo, no collage grid, "
    "no people, and no extra props unless the product itself is decor."
)


@dataclass
class ProductImageSpec:
    product_id: int
    product_detail: str
    image_filename: str
    image_prompt: str


def extract_folder_id(folder_url: str) -> str:
    match = re.search(r"/folders/([A-Za-z0-9_-]+)", folder_url)
    if match:
        return match.group(1)
    if re.fullmatch(r"[A-Za-z0-9_-]{20,}", folder_url):
        return folder_url
    raise ValueError(f"Unable to extract Google Drive folder ID from: {folder_url}")


def load_products(workbook_path: Path) -> list[ProductImageSpec]:
    workbook = openpyxl.load_workbook(workbook_path, read_only=True, data_only=True)
    worksheet = workbook[workbook.sheetnames[0]]
    rows = worksheet.iter_rows(values_only=True)
    header = next(rows)
    header_map = {str(name): index for index, name in enumerate(header)}

    required_columns = [
        "product_id",
        "product_detail",
        "image_filename",
        "image_prompt",
    ]
    missing = [name for name in required_columns if name not in header_map]
    if missing:
        raise ValueError(f"Missing required columns: {', '.join(missing)}")

    products: list[ProductImageSpec] = []
    for row in rows:
        image_filename = row[header_map["image_filename"]]
        image_prompt = row[header_map["image_prompt"]]
        if not image_filename or not image_prompt:
            continue
        products.append(
            ProductImageSpec(
                product_id=int(row[header_map["product_id"]]),
                product_detail=str(row[header_map["product_detail"]]),
                image_filename=str(image_filename),
                image_prompt=str(image_prompt).strip(),
            )
        )
    return products


def fetch_image_bytes(
    session: requests.Session,
    prompt: str,
    product_id: int,
    *,
    model: str,
    width: int,
    height: int,
    retries: int,
    pause_seconds: float,
) -> bytes:
    encoded_prompt = quote(prompt, safe="")
    url = f"{API_BASE}{encoded_prompt}"
    params = {
        "model": model,
        "width": width,
        "height": height,
        "seed": product_id,
        "nologo": "true",
        "safe": "true",
    }

    last_error = "Unknown error"
    for attempt in range(1, retries + 1):
        try:
            response = session.get(url, params=params, timeout=240)
            content_type = (response.headers.get("content-type") or "").lower()
            if response.ok and content_type.startswith("image/"):
                return response.content
            last_error = f"HTTP {response.status_code}: {response.text[:300]}"
        except requests.RequestException as exc:
            last_error = f"{type(exc).__name__}: {exc}"

        if attempt < retries:
            time.sleep(pause_seconds * attempt)

    raise RuntimeError(last_error)


def build_prompt(product: ProductImageSpec) -> str:
    return f"{product.image_prompt} {REFERENCE_STYLE}"


def list_drive_folder_files(folder_url: str, session: requests.Session) -> dict[str, str]:
    folder_id = extract_folder_id(folder_url)
    embedded_url = f"https://drive.google.com/embeddedfolderview?id={folder_id}#list"
    response = session.get(embedded_url, timeout=120)
    response.raise_for_status()
    html = response.text

    matches = re.findall(
        r'<div class="flip-entry" id="entry-([A-Za-z0-9_-]+)".*?<div class="flip-entry-title">([^<]+)</div>',
        html,
        flags=re.DOTALL,
    )
    return {
        title.strip(): f"https://drive.google.com/uc?export=view&id={file_id}"
        for file_id, title in matches
    }


def update_workbook_image_urls(
    workbook_path: Path,
    products: list[ProductImageSpec],
    folder_url: str,
    session: requests.Session,
    *,
    sync_wait_seconds: int,
    sync_poll_seconds: int,
) -> None:
    expected_filenames = {product.image_filename for product in products}
    deadline = time.time() + sync_wait_seconds
    drive_links: dict[str, str] = {}

    while time.time() < deadline:
        drive_links = list_drive_folder_files(folder_url, session)
        available = expected_filenames.intersection(drive_links)
        print(f"Drive folder synced: {len(available)}/{len(expected_filenames)} expected files visible")
        if expected_filenames.issubset(drive_links):
            break
        time.sleep(sync_poll_seconds)

    missing = sorted(expected_filenames.difference(drive_links))
    if missing:
        raise RuntimeError(
            "Drive sync did not expose all files before timeout. Missing examples: "
            + ", ".join(missing[:10])
        )

    workbook = openpyxl.load_workbook(workbook_path)
    worksheet = workbook[workbook.sheetnames[0]]
    header = [cell.value for cell in next(worksheet.iter_rows(min_row=1, max_row=1))]
    header_map = {str(name): index for index, name in enumerate(header)}
    image_filename_index = header_map["image_filename"] + 1
    image_url_index = header_map["image_url"] + 1

    for row_index in range(2, worksheet.max_row + 1):
        filename = worksheet.cell(row=row_index, column=image_filename_index).value
        if filename in drive_links:
            worksheet.cell(row=row_index, column=image_url_index).value = drive_links[filename]

    workbook.save(workbook_path)
    print(f"Workbook updated: {workbook_path}")


def save_png(image_bytes: bytes, output_path: Path) -> None:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    with Image.open(io.BytesIO(image_bytes)) as image:
        image.convert("RGB").save(output_path, format="PNG", optimize=True)


def iter_pending_products(
    products: Iterable[ProductImageSpec],
    output_dir: Path,
    overwrite: bool,
) -> Iterable[ProductImageSpec]:
    for product in products:
        output_path = output_dir / product.image_filename
        if overwrite or not output_path.exists():
            yield product


def main() -> int:
    parser = argparse.ArgumentParser(description="Generate product images from workbook prompts.")
    parser.add_argument(
        "--workbook",
        default="Product_PowerBI_Images.xlsx",
        help="Path to the workbook containing image prompts.",
    )
    parser.add_argument(
        "--output-dir",
        default="img",
        help="Folder where generated product images will be saved.",
    )
    parser.add_argument("--model", default="flux", help="Pollinations image model name.")
    parser.add_argument("--width", type=int, default=1024, help="Requested image width.")
    parser.add_argument("--height", type=int, default=1024, help="Requested image height.")
    parser.add_argument("--retries", type=int, default=3, help="Retries per image.")
    parser.add_argument(
        "--pause-seconds",
        type=float,
        default=2.0,
        help="Base pause between retries and requests.",
    )
    parser.add_argument(
        "--overwrite",
        action="store_true",
        help="Overwrite images that already exist in the output folder.",
    )
    parser.add_argument(
        "--start-at",
        type=int,
        default=1,
        help="1-based index of the product row to start from after filtering pending items.",
    )
    parser.add_argument(
        "--count",
        type=int,
        default=0,
        help="Number of pending items to generate. Use 0 to process all pending items.",
    )
    parser.add_argument(
        "--drive-folder-url",
        default="",
        help="Optional shared Google Drive folder URL used to write public image links back into the workbook.",
    )
    parser.add_argument(
        "--sync-wait-seconds",
        type=int,
        default=900,
        help="How long to wait for Google Drive sync before updating workbook URLs.",
    )
    parser.add_argument(
        "--sync-poll-seconds",
        type=int,
        default=15,
        help="Polling interval while waiting for Google Drive sync.",
    )
    args = parser.parse_args()

    workbook_path = Path(args.workbook).resolve()
    output_dir = Path(args.output_dir).resolve()
    products = load_products(workbook_path)
    pending_products = list(iter_pending_products(products, output_dir, args.overwrite))
    if args.start_at < 1:
        raise ValueError("--start-at must be 1 or greater")
    start_index = args.start_at - 1
    if args.count > 0:
        pending_products = pending_products[start_index : start_index + args.count]
    else:
        pending_products = pending_products[start_index:]

    print(f"Workbook: {workbook_path}")
    print(f"Output dir: {output_dir}")
    print(f"Products in sheet: {len(products)}")
    print(f"Products to generate: {len(pending_products)}")

    session = requests.Session()
    session.headers["User-Agent"] = "Codex Product Image Generator/1.0"

    failures: list[str] = []
    for index, product in enumerate(pending_products, start=1):
        output_path = output_dir / product.image_filename
        print(f"[{index}/{len(pending_products)}] Generating {product.image_filename} for {product.product_detail}")
        try:
            image_bytes = fetch_image_bytes(
                session,
                build_prompt(product),
                product.product_id,
                model=args.model,
                width=args.width,
                height=args.height,
                retries=args.retries,
                pause_seconds=args.pause_seconds,
            )
            save_png(image_bytes, output_path)
            print(f"Saved: {output_path}")
        except Exception as exc:  # noqa: BLE001
            failures.append(f"{product.image_filename}: {exc}")
            print(f"Failed: {product.image_filename} -> {exc}")
        time.sleep(args.pause_seconds)

    print("")
    print(f"Completed with {len(failures)} failure(s).")
    for failure in failures:
        print(failure)

    if not failures and args.drive_folder_url:
        update_workbook_image_urls(
            workbook_path,
            products,
            args.drive_folder_url,
            session,
            sync_wait_seconds=args.sync_wait_seconds,
            sync_poll_seconds=args.sync_poll_seconds,
        )

    return 1 if failures else 0


if __name__ == "__main__":
    raise SystemExit(main())
